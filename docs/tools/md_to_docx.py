#!/usr/bin/env python3
"""Convert the zbsniff markdown user docs to .docx.

Image handling:
  images/<name>.svg  → the Pillow-rendered docs/tools/_figs/<name>.png (concept diagrams)
  images/<name>.png  → docs/images/<name>.png used as-is (real screenshots you drop in)

Run `render_figures.py` first (or use build.sh), then this. Deps: python-docx, Pillow.
"""
import os, re, sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

TOOLS = os.path.dirname(os.path.abspath(__file__))
DOCS = os.path.dirname(TOOLS)
FIGS = os.path.join(TOOLS, "_figs")

def image_path(rel):
    """Map a markdown image ref to a real PNG on disk (or None if missing)."""
    if rel.endswith(".svg"):
        p = os.path.join(FIGS, os.path.basename(rel).replace(".svg", ".png"))
    else:
        p = os.path.join(DOCS, rel)
    return p if os.path.exists(p) else None

INLINE = re.compile(r'(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))')

def add_runs(paragraph, text):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = paragraph.add_run(part[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)
        elif part.startswith("["):
            m = re.match(r'\[([^\]]+)\]\(([^)]+)\)', part)
            label, url = m.group(1), m.group(2)
            paragraph.add_run(label if not url.startswith("http") else f"{label} ({url})")
        else:
            paragraph.add_run(part)

def flush_table(doc, rows):
    header = rows[0]
    sep = len(rows) > 1 and set("".join(rows[1]).replace("|", "").strip()) <= set("-: ")
    body = rows[2:] if sep else rows[1:]
    t = doc.add_table(rows=1, cols=len(header)); t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, c in enumerate(header):
        p = t.rows[0].cells[i].paragraphs[0]; add_runs(p, c.strip())
        for run in p.runs: run.bold = True
    for row in body:
        cells = t.add_row().cells
        for i, c in enumerate(row[:len(header)]):
            add_runs(cells[i].paragraphs[0], c.strip())

def convert(md_path, out_path):
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
    lines = open(md_path).read().split("\n")
    i, n, tbl = 0, len(lines), []
    def flush():
        nonlocal tbl
        if tbl: flush_table(doc, tbl); tbl = []
    while i < n:
        line = lines[i]
        if line.startswith("```"):
            flush(); i += 1; code = []
            while i < n and not lines[i].startswith("```"): code.append(lines[i]); i += 1
            i += 1
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.2)
            r = p.add_run("\n".join(code)); r.font.name = "Consolas"; r.font.size = Pt(9)
            continue
        if line.strip().startswith("|") and line.strip().endswith("|"):
            tbl.append(line.strip().strip("|").split("|")); i += 1; continue
        flush()
        m = re.match(r'!\[([^\]]*)\]\(([^)]+)\)', line.strip())
        if m:
            img = image_path(m.group(2))
            if img:
                p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(img, width=Inches(6.3))
            else:
                doc.add_paragraph(f"[figure not found: {m.group(2)}]")
            i += 1; continue
        if line.startswith("#"):
            lvl = len(line) - len(line.lstrip("#"))
            doc.add_heading(line.lstrip("#").strip(), level=0 if lvl == 1 else min(lvl-1, 4))
            i += 1; continue
        if line.strip() in ("---", "***"):
            doc.add_paragraph().add_run("─" * 40).font.color.rgb = RGBColor(0xBB, 0xBB, 0xBB)
            i += 1; continue
        if line.startswith(">"):
            add_runs(doc.add_paragraph(style="Intense Quote"), line.lstrip(">").strip()); i += 1; continue
        mb = re.match(r'^\s*-\s+(.*)', line)
        if mb: add_runs(doc.add_paragraph(style="List Bullet"), mb.group(1)); i += 1; continue
        mn = re.match(r'^\s*\d+\.\s+(.*)', line)
        if mn: add_runs(doc.add_paragraph(style="List Number"), mn.group(1)); i += 1; continue
        if not line.strip(): i += 1; continue
        add_runs(doc.add_paragraph(), line); i += 1
    flush()
    doc.save(out_path); print("wrote", out_path)

if __name__ == "__main__":
    convert(os.path.join(DOCS, "quickstart.md"), os.path.join(DOCS, "quickstart.docx"))
    convert(os.path.join(DOCS, "user-guide.md"), os.path.join(DOCS, "user-guide.docx"))
